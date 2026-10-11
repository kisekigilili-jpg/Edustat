import io
import os
import re
import tempfile
import unittest
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook
from fastapi.testclient import TestClient

from app import database


os.environ["INITIAL_TEACHER_USERNAME"] = "teacher"
os.environ["INITIAL_TEACHER_PASSWORD"] = "correct-horse-battery-staple"
os.environ["SESSION_SECRET"] = "test-session-secret-that-is-long-enough"

from app.main import app


class AppTests(unittest.TestCase):
    api_auth = ("teacher", "correct-horse-battery-staple")

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.original_db_path = database.DB_PATH
        database.DB_PATH = Path(self.temp_dir.name) / "test.db"
        self.client = TestClient(app)
        self.client.__enter__()
        login_page = self.client.get("/login")
        token = re.search(r'name="csrf_token" value="([^"]+)"', login_page.text).group(1)
        response = self.client.post(
            "/login",
            data={
                "csrf_token": token,
                "username": "teacher",
                "password": "correct-horse-battery-staple",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)

    def tearDown(self):
        self.client.__exit__(None, None, None)
        database.DB_PATH = self.original_db_path
        self.temp_dir.cleanup()

    def csrf_token(self, path="/"):
        response = self.client.get(path)
        self.assertEqual(response.status_code, 200)
        return re.search(
            r'name="csrf_token" value="([^"]+)"', response.text
        ).group(1)

    def test_students_assessments_reports_and_excel_export(self):
        token = self.csrf_token()
        response = self.client.post(
            "/students",
            data={
                "csrf_token": token,
                "full_name": "=2+2",
                "grade": "Grade 10",
                "student_number": "STU-101",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)

        dashboard = self.client.get("/")
        student_id = re.search(r'href="/students/(\d+)"', dashboard.text).group(1)
        detail_path = f"/students/{student_id}"
        token = self.csrf_token(detail_path)
        response = self.client.post(
            f"{detail_path}/assessments",
            data={
                "csrf_token": token,
                "category": "Test",
                "title": "Fractions",
                "score": "8",
                "max_score": "10",
                "assessed_on": "2026-10-10",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)

        report = self.client.get("/reports?search=STU-101")
        self.assertIn("80.0%", report.text)
        export = self.client.get("/export.xlsx")
        self.assertEqual(export.status_code, 200)
        workbook = load_workbook(filename=io.BytesIO(export.content))
        self.assertEqual(workbook.sheetnames, ["Student Summary", "Assessments"])
        self.assertEqual(workbook["Student Summary"]["C2"].value, "'=2+2")
        self.assertEqual(workbook["Student Summary"]["H2"].value, 80)
        self.assertEqual(workbook["Assessments"]["F2"].value, "Fractions")

    def test_invalid_score_and_invalid_csrf_are_rejected(self):
        token = self.csrf_token()
        response = self.client.post(
            "/students",
            data={
                "csrf_token": token,
                "full_name": "Alex Morgan",
                "grade": "Grade 10",
                "student_number": "",
            },
        )
        self.assertEqual(response.status_code, 200)
        student_id = re.search(r'href="/students/(\d+)"', response.text).group(1)
        detail_path = f"/students/{student_id}"
        token = self.csrf_token(detail_path)
        invalid_score = self.client.post(
            f"{detail_path}/assessments",
            data={
                "csrf_token": token,
                "category": "Exam",
                "title": "Final exam",
                "score": "12",
                "max_score": "10",
                "assessed_on": "2026-10-10",
            },
        )
        self.assertEqual(invalid_score.status_code, 422)
        invalid_csrf = self.client.post(
            f"{detail_path}/assessments",
            data={
                "csrf_token": "invalid",
                "category": "Exam",
                "title": "Final exam",
                "score": "8",
                "max_score": "10",
                "assessed_on": "2026-10-10",
            },
        )
        self.assertEqual(invalid_csrf.status_code, 403)

    def test_password_can_be_changed_and_used_to_sign_in(self):
        token = self.csrf_token("/settings/password")
        response = self.client.post(
            "/settings/password",
            data={
                "csrf_token": token,
                "current_password": "correct-horse-battery-staple",
                "new_password": "another-correct-password",
                "confirm_password": "another-correct-password",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)
        self.assertTrue(urlparse(response.headers["location"]).path.endswith("/password"))

        token = self.csrf_token("/")
        self.client.post("/logout", data={"csrf_token": token}, follow_redirects=False)
        token = self.csrf_token("/login")
        response = self.client.post(
            "/login",
            data={
                "csrf_token": token,
                "username": "teacher",
                "password": "another-correct-password",
            },
            follow_redirects=False,
        )
        self.assertEqual(response.status_code, 303)

    def test_api_registration_marks_and_updates(self):
        student_response = self.client.post(
            "/api/v1/students",
            auth=self.api_auth,
            json={
                "full_name": "  Alex Morgan ",
                "grade": " Grade 10 ",
                "student_number": "STU-API-1",
            },
        )
        self.assertEqual(student_response.status_code, 201)
        student = student_response.json()
        student_id = student["id"]
        self.assertEqual(student["full_name"], "Alex Morgan")
        self.assertEqual(student["grade"], "Grade 10")

        marks_url = f"/api/v1/students/{student_id}/marks"
        mark_response = self.client.post(
            marks_url,
            auth=self.api_auth,
            json={
                "category": "Assignment",
                "title": "  Homework 1 ",
                "score": 18,
                "max_score": 20,
                "assessed_on": "2026-10-10",
            },
        )
        self.assertEqual(mark_response.status_code, 201)
        mark = mark_response.json()
        self.assertEqual(mark["title"], "Homework 1")
        self.assertEqual(mark["percentage"], 90)
        mark_id = mark["id"]

        detail = self.client.get(
            f"/api/v1/students/{student_id}", auth=self.api_auth
        ).json()
        self.assertEqual(detail["summary"]["overall_mark"], 90)
        self.assertEqual(detail["summary"]["assessment_count"], 1)
        self.assertEqual(detail["marks"][0]["id"], mark_id)

        marks = self.client.get(marks_url, auth=self.api_auth)
        self.assertEqual(marks.status_code, 200)
        self.assertEqual(len(marks.json()), 1)

        updated_mark = self.client.patch(
            f"{marks_url}/{mark_id}",
            auth=self.api_auth,
            json={"score": 19},
        )
        self.assertEqual(updated_mark.status_code, 200)
        self.assertEqual(updated_mark.json()["percentage"], 95)

        updated_student = self.client.patch(
            f"/api/v1/students/{student_id}",
            auth=self.api_auth,
            json={"grade": "Grade 11"},
        )
        self.assertEqual(updated_student.status_code, 200)
        self.assertEqual(updated_student.json()["grade"], "Grade 11")

        searched = self.client.get(
            "/api/v1/students",
            auth=self.api_auth,
            params={"search": "STU-API-1", "limit": 10},
        )
        self.assertEqual(searched.status_code, 200)
        self.assertEqual(searched.json()["total"], 1)
        self.assertEqual(searched.json()["items"][0]["id"], student_id)

        deleted_mark = self.client.delete(
            f"{marks_url}/{mark_id}", auth=self.api_auth
        )
        self.assertEqual(deleted_mark.status_code, 204)
        deleted_student = self.client.delete(
            f"/api/v1/students/{student_id}", auth=self.api_auth
        )
        self.assertEqual(deleted_student.status_code, 204)
        missing_student = self.client.get(
            f"/api/v1/students/{student_id}", auth=self.api_auth
        )
        self.assertEqual(missing_student.status_code, 404)

    def test_api_auth_validation_duplicate_numbers_and_openapi_tags(self):
        unauthorized = self.client.get("/api/v1/students")
        self.assertEqual(unauthorized.status_code, 401)
        self.assertEqual(unauthorized.json()["detail"], "Teacher authentication is required.")
        invalid_credentials = self.client.get(
            "/api/v1/students", auth=("teacher", "not-the-password")
        )
        self.assertEqual(invalid_credentials.status_code, 401)

        docs = self.client.get("/openapi.json")
        self.assertEqual(docs.status_code, 200)
        openapi = docs.json()
        self.assertEqual(
            {tag["name"] for tag in openapi["tags"]},
            {"Students", "Marks"},
        )
        self.assertIn(
            "HTTPBasic",
            openapi["paths"]["/api/v1/students"]["get"]["security"][0],
        )

        payload = {
            "full_name": "Jamie Lee",
            "grade": "Grade 9",
            "student_number": "STU-DUPLICATE",
        }
        created = self.client.post(
            "/api/v1/students", auth=self.api_auth, json=payload
        )
        self.assertEqual(created.status_code, 201)
        duplicate = self.client.post(
            "/api/v1/students", auth=self.api_auth, json=payload
        )
        self.assertEqual(duplicate.status_code, 409)
        blank_name = self.client.post(
            "/api/v1/students",
            auth=self.api_auth,
            json={"full_name": " ", "grade": "Grade 9"},
        )
        self.assertEqual(blank_name.status_code, 422)

        student_id = created.json()["id"]
        invalid_mark = self.client.post(
            f"/api/v1/students/{student_id}/marks",
            auth=self.api_auth,
            json={
                "category": "Exam",
                "title": "Final",
                "score": 11,
                "max_score": 10,
                "assessed_on": "2026-10-10",
            },
        )
        self.assertEqual(invalid_mark.status_code, 422)


if __name__ == "__main__":
    unittest.main()
