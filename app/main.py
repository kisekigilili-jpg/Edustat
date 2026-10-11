import io
import logging
import math
import os
import secrets
import sqlite3
from datetime import date
from pathlib import Path
from urllib.parse import urlparse

from fastapi import FastAPI, Form, HTTPException, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from starlette.middleware.sessions import SessionMiddleware

from app.api.routers.marks import router as marks_api_router
from app.api.routers.students import router as students_api_router
from app.database import connect, initialize_database
from app.grades import CATEGORIES, student_summaries, summarize_assessments
from app.security import hash_password, verify_password


BASE_DIR = Path(__file__).resolve().parent
logger = logging.getLogger("student_marks")

app = FastAPI(
    title="Student Marks API",
    description=(
        "Register students and record assignment, test, and exam marks. "
        "API endpoints require teacher HTTP Basic authentication."
    ),
    version="1.0.0",
    openapi_tags=[
        {
            "name": "Students",
            "description": "Student registration, lookup, search, and roster management.",
        },
        {
            "name": "Marks",
            "description": "Record, review, update, and remove assignment, test, and exam marks.",
        },
    ],
)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=BASE_DIR / "templates")
app.include_router(students_api_router, prefix="/api/v1")
app.include_router(marks_api_router, prefix="/api/v1")


def csrf_token(request: Request) -> str:
    token = request.session.get("csrf_token")
    if token is None:
        token = secrets.token_urlsafe(32)
        request.session["csrf_token"] = token
    return token


def check_csrf(request: Request, submitted_token: str):
    expected_token = request.session.get("csrf_token", "")
    if not expected_token or not secrets.compare_digest(expected_token, submitted_token):
        raise HTTPException(status_code=403, detail="Invalid form token. Please reload and try again.")


def render(request: Request, template_name: str, **context):
    context.setdefault("csrf_token", csrf_token(request))
    context.setdefault("username", request.session.get("username"))
    return templates.TemplateResponse(
        request=request,
        name=template_name,
        context=context,
    )


@app.on_event("startup")
def startup():
    initialize_database()
    with connect() as connection:
        teacher_exists = connection.execute("SELECT 1 FROM teachers LIMIT 1").fetchone()
    if teacher_exists:
        return

    username = os.getenv("INITIAL_TEACHER_USERNAME", "teacher").strip()
    password = os.getenv("INITIAL_TEACHER_PASSWORD")
    generated_password = password is None
    if generated_password:
        password = secrets.token_urlsafe(15)
    if not username or len(username) > 80 or len(password) < 12:
        raise RuntimeError(
            "Set INITIAL_TEACHER_USERNAME and INITIAL_TEACHER_PASSWORD "
            "(the password must be at least 12 characters)."
        )
    salt, password_hash = hash_password(password)
    with connect() as connection:
        connection.execute(
            "INSERT INTO teachers (username, password_salt, password_hash) VALUES (?, ?, ?)",
            (username, salt, password_hash),
        )
    if generated_password:
        logger.warning(
            "Created the first teacher account. Sign-in username: %s; "
            "one-time initial password: %s. Change this password after signing in.",
            username,
            password,
        )


@app.middleware("http")
async def require_teacher(request: Request, call_next):
    if request.url.path.startswith("/api/") or request.url.path in {
        "/docs",
        "/redoc",
        "/openapi.json",
    }:
        return await call_next(request)
    public_path = (
        request.url.path == "/login"
        or request.url.path.startswith("/static/")
    )
    if not public_path and not request.session.get("teacher_id"):
        return RedirectResponse(url="/login", status_code=303)
    return await call_next(request)


app.add_middleware(
    SessionMiddleware,
    secret_key=os.getenv("SESSION_SECRET") or secrets.token_urlsafe(32),
    same_site="lax",
    https_only=os.getenv("COOKIE_SECURE", "false").lower() == "true",
)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, error: str = ""):
    return render(request, "login.html", error=error, next_url="/")


@app.post("/login")
def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    submitted_csrf: str = Form(..., alias="csrf_token"),
    next_url: str = Form("/", alias="next"),
):
    check_csrf(request, submitted_csrf)
    with connect() as connection:
        teacher = connection.execute(
            "SELECT * FROM teachers WHERE username = ?", (username.strip(),)
        ).fetchone()

    if teacher is None or not verify_password(
        password, teacher["password_salt"], teacher["password_hash"]
    ):
        return render(
            request,
            "login.html",
            error="The username or password is incorrect.",
            next_url="/",
        )

    parsed_next = urlparse(next_url)
    safe_next = (
        next_url
        if parsed_next.path.startswith("/") and not parsed_next.netloc
        else "/"
    )
    request.session.clear()
    request.session["teacher_id"] = teacher["id"]
    request.session["username"] = teacher["username"]
    request.session["csrf_token"] = secrets.token_urlsafe(32)
    return RedirectResponse(url=safe_next, status_code=303)


@app.post("/logout")
def logout(
    request: Request,
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    request.session.clear()
    return RedirectResponse(url="/login", status_code=303)


@app.get("/", response_class=HTMLResponse)
def dashboard(request: Request):
    students = student_summaries()
    assessment_count = sum(
        student["summary"]["assessment_count"] for student in students
    )
    return render(
        request,
        "dashboard.html",
        students=students,
        student_count=len(students),
        assessment_count=assessment_count,
    )


@app.post("/students")
def create_student(
    request: Request,
    full_name: str = Form(...),
    grade: str = Form(...),
    student_number: str = Form(""),
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    full_name = full_name.strip()
    grade = grade.strip()
    student_number = student_number.strip()
    if not full_name or len(full_name) > 120:
        raise HTTPException(status_code=422, detail="Student name must be 1-120 characters.")
    if not grade or len(grade) > 60:
        raise HTTPException(status_code=422, detail="Grade or class must be 1-60 characters.")
    if student_number and len(student_number) > 40:
        raise HTTPException(status_code=422, detail="Student number cannot exceed 40 characters.")

    try:
        with connect() as connection:
            connection.execute(
                "INSERT INTO students (student_number, full_name, grade) VALUES (?, ?, ?)",
                (student_number, full_name, grade),
            )
    except sqlite3.IntegrityError as exc:
        raise HTTPException(
            status_code=409,
            detail="That student number is already in use.",
        ) from exc
    return RedirectResponse(url="/", status_code=303)


@app.get("/students/{student_id}", response_class=HTMLResponse)
def student_detail(request: Request, student_id: int):
    with connect() as connection:
        student = connection.execute(
            "SELECT * FROM students WHERE id = ?", (student_id,)
        ).fetchone()
        if student is None:
            raise HTTPException(status_code=404, detail="Student not found.")
        assessments = connection.execute(
            """
            SELECT * FROM assessments
            WHERE student_id = ?
            ORDER BY assessed_on DESC, id DESC
            """,
            (student_id,),
        ).fetchall()
    return render(
        request,
        "student_detail.html",
        student=student,
        assessments=assessments,
        summary=summarize_assessments(assessments),
        categories=CATEGORIES,
        today=date.today().isoformat(),
    )


@app.post("/students/{student_id}/assessments")
def create_assessment(
    request: Request,
    student_id: int,
    category: str = Form(...),
    title: str = Form(...),
    score: float = Form(...),
    max_score: float = Form(...),
    assessed_on: str = Form(...),
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    if category not in CATEGORIES:
        raise HTTPException(status_code=422, detail="Choose Assignment, Test, or Exam.")
    title = title.strip()
    if not title or len(title) > 120:
        raise HTTPException(status_code=422, detail="Assessment title must be 1-120 characters.")
    if (
        not math.isfinite(score)
        or not math.isfinite(max_score)
        or max_score <= 0
        or score < 0
        or score > max_score
    ):
        raise HTTPException(
            status_code=422,
            detail="Score must be between 0 and the maximum score, which must be greater than 0.",
        )
    try:
        date.fromisoformat(assessed_on)
    except ValueError as exc:
        raise HTTPException(status_code=422, detail="Enter a valid assessment date.") from exc

    with connect() as connection:
        student_exists = connection.execute(
            "SELECT 1 FROM students WHERE id = ?", (student_id,)
        ).fetchone()
        if student_exists is None:
            raise HTTPException(status_code=404, detail="Student not found.")
        connection.execute(
            """
            INSERT INTO assessments (student_id, category, title, score, max_score, assessed_on)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (student_id, category, title, score, max_score, assessed_on),
        )
    return RedirectResponse(url=f"/students/{student_id}", status_code=303)


@app.post("/assessments/{assessment_id}/delete")
def delete_assessment(
    request: Request,
    assessment_id: int,
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    with connect() as connection:
        assessment = connection.execute(
            "SELECT student_id FROM assessments WHERE id = ?", (assessment_id,)
        ).fetchone()
        if assessment is None:
            raise HTTPException(status_code=404, detail="Assessment not found.")
        connection.execute("DELETE FROM assessments WHERE id = ?", (assessment_id,))
    return RedirectResponse(url=f"/students/{assessment['student_id']}", status_code=303)


@app.post("/students/{student_id}/delete")
def delete_student(
    request: Request,
    student_id: int,
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    with connect() as connection:
        cursor = connection.execute("DELETE FROM students WHERE id = ?", (student_id,))
        if cursor.rowcount == 0:
            raise HTTPException(status_code=404, detail="Student not found.")
    return RedirectResponse(url="/", status_code=303)


@app.get("/reports", response_class=HTMLResponse)
def reports(request: Request, search: str = ""):
    search = search.strip()[:100]
    return render(
        request,
        "reports.html",
        students=student_summaries(search),
        search=search,
    )


def excel_text(value):
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


@app.get("/export.xlsx")
def export_excel():
    students = student_summaries()
    with connect() as connection:
        assessments = connection.execute(
            """
            SELECT a.*, s.full_name, s.student_number, s.grade
            FROM assessments a
            JOIN students s ON s.id = a.student_id
            ORDER BY s.full_name COLLATE NOCASE, a.assessed_on, a.id
            """
        ).fetchall()

    workbook = Workbook()
    summary_sheet = workbook.active or workbook.create_sheet()
    summary_sheet.title = "Student Summary"
    summary_sheet.append(
        [
            "Student ID",
            "Student Number",
            "Full Name",
            "Grade / Class",
            "Assignments Average (%)",
            "Tests Average (%)",
            "Exams Average (%)",
            "Overall (%)",
            "Assessment Count",
        ]
    )
    for student in students:
        summary = student["summary"]
        summary_sheet.append(
            [
                student["id"],
                excel_text(student["student_number"]),
                excel_text(student["full_name"]),
                excel_text(student["grade"]),
                *[
                    round(summary["categories"][category], 1)
                    if summary["categories"][category] is not None
                    else None
                    for category in CATEGORIES
                ],
                round(summary["overall"], 1) if summary["overall"] is not None else None,
                summary["assessment_count"],
            ]
        )

    detail_sheet = workbook.create_sheet("Assessments")
    detail_sheet.append(
        [
            "Student ID",
            "Student Number",
            "Full Name",
            "Grade / Class",
            "Category",
            "Assessment",
            "Score",
            "Maximum Score",
            "Result (%)",
            "Date",
        ]
    )
    for assessment in assessments:
        detail_sheet.append(
            [
                assessment["student_id"],
                excel_text(assessment["student_number"]),
                excel_text(assessment["full_name"]),
                excel_text(assessment["grade"]),
                assessment["category"],
                excel_text(assessment["title"]),
                assessment["score"],
                assessment["max_score"],
                round(assessment["score"] / assessment["max_score"] * 100, 1),
                assessment["assessed_on"],
            ]
        )

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24463A")
        for column_index, column in enumerate(sheet.columns, start=1):
            width = min(max(max(len(str(cell.value or "")) for cell in column) + 2, 12), 36)
            sheet.column_dimensions[get_column_letter(column_index)].width = width

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="student-marks.xlsx"'},
    )


@app.get("/settings/password", response_class=HTMLResponse)
def password_settings(request: Request, message: str = ""):
    return render(request, "password.html", message=message, error="")


@app.post("/settings/password")
def change_password(
    request: Request,
    current_password: str = Form(...),
    new_password: str = Form(...),
    confirm_password: str = Form(...),
    submitted_csrf: str = Form(..., alias="csrf_token"),
):
    check_csrf(request, submitted_csrf)
    if len(new_password) < 12:
        return render(
            request,
            "password.html",
            message="",
            error="Use at least 12 characters for the new password.",
        )
    if new_password != confirm_password:
        return render(
            request,
            "password.html",
            message="",
            error="The new password and confirmation do not match.",
        )

    with connect() as connection:
        teacher = connection.execute(
            "SELECT * FROM teachers WHERE id = ?",
            (request.session["teacher_id"],),
        ).fetchone()
        if teacher is None or not verify_password(
            current_password, teacher["password_salt"], teacher["password_hash"]
        ):
            return render(
                request,
                "password.html",
                message="",
                error="The current password is incorrect.",
            )
        salt, password_hash = hash_password(new_password)
        connection.execute(
            "UPDATE teachers SET password_salt = ?, password_hash = ? WHERE id = ?",
            (salt, password_hash, teacher["id"]),
        )

    return RedirectResponse(
        url="/settings/password?message=Password%20updated%20successfully.",
        status_code=303,
    )
